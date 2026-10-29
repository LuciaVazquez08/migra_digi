"""Fase 1: inventario de solo lectura del origen y lectura del Excel de referencia."""
from __future__ import annotations

import logging
import os
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .archivos import largo
from .modelo import Entrada
from .reglas import RE_CONTRATO, RE_SDS, categoria_cliente, es_ignorado

log = logging.getLogger(__name__)


def _nfc(texto: str) -> str:
    return unicodedata.normalize("NFC", texto)


def listar_raiz(origen: Path) -> list[tuple[str, Path, bool]]:
    """Elementos de primer nivel del origen ordenados por nombre: (nombre, ruta, es_dir)."""
    with os.scandir(largo(origen)) as it:
        elementos = [(_nfc(e.name), origen / e.name, e.is_dir(follow_symlinks=False)) for e in it]
    return sorted(elementos, key=lambda x: x[0])


def escanear_cliente(ruta: Path) -> list[Entrada]:
    """Recorre recursivamente la carpeta de un cliente sin modificar nada."""
    entradas: list[Entrada] = []

    def recorrer(actual: Path, rel: tuple[str, ...]) -> None:
        with os.scandir(largo(actual)) as it:
            hijos = sorted(it, key=lambda e: e.name)
        for e in hijos:
            r = rel + (_nfc(e.name),)
            p = actual / e.name
            if e.is_symlink():
                # Un enlace simbólico/junction no es documentación verificable.
                entradas.append(Entrada(r, False, -1, p))
                continue
            if e.is_dir():
                entradas.append(Entrada(r, True, 0, p))
                recorrer(p, r)
            else:
                entradas.append(Entrada(r, False, e.stat().st_size, p))

    recorrer(ruta, ())
    return entradas


@dataclass
class Referencia:
    """Contenido del Excel: por cada carpeta de nivel 1, el conjunto de rutas relativas."""
    arbol: dict[str, set[tuple[str, ...]]] = field(default_factory=dict)
    # Índices globales para detectar SDS/contratos repetidos en varios clientes
    sds_por_cliente: dict[str, set[str]] = field(default_factory=lambda: defaultdict(set))
    contratos_por_cliente: dict[str, set[str]] = field(default_factory=lambda: defaultdict(set))

    def clientes_de_sds(self, sds: str) -> set[str]:
        return self.sds_por_cliente.get(sds, set())

    def clientes_de_contrato(self, contrato: str) -> set[str]:
        return self.contratos_por_cliente.get(contrato, set())


def _celda(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = _nfc(str(valor))
    return texto if texto != "" else None


def leer_referencia(excel: Path, hoja: str) -> Referencia:
    """Lee la hoja de árbol (columnas ``Nivel 1``..``Nivel N``) en streaming."""
    log.info("Leyendo Excel de referencia %s (hoja %s)", excel, hoja)
    wb = load_workbook(excel, read_only=True, data_only=True)
    try:
        if hoja not in wb.sheetnames:
            raise ValueError(f"El Excel no tiene la hoja '{hoja}' (hojas: {wb.sheetnames})")
        ws = wb[hoja]
        filas = ws.iter_rows(values_only=True)
        encabezado = next(filas)
        columnas = [i for i, c in enumerate(encabezado) if str(c or "").startswith("Nivel")]
        if not columnas or columnas[0] != 0:
            raise ValueError(f"Encabezado inesperado en '{hoja}': {encabezado}")
        ref = Referencia()
        for fila in filas:
            niveles = []
            for i in columnas:
                v = _celda(fila[i]) if i < len(fila) else None
                if v is None:
                    break
                niveles.append(v)
            if not niveles:
                continue
            cliente, rel = niveles[0], tuple(niveles[1:])
            ref.arbol.setdefault(cliente, set())
            if rel:
                ref.arbol[cliente].add(rel)
                categoria = categoria_cliente(rel[0])[0]
                for nombre in rel[1:]:
                    if categoria == "movil" and RE_SDS.fullmatch(nombre):
                        ref.sds_por_cliente[nombre].add(cliente)
                    elif categoria == "contratos" and RE_CONTRATO.fullmatch(nombre):
                        ref.contratos_por_cliente[nombre].add(cliente)
    finally:
        wb.close()
    log.info("Excel leído: %d carpetas de nivel 1", len(ref.arbol))
    return ref


def comparar_con_referencia(entradas: list[Entrada], esperado: set[tuple[str, ...]]
                            ) -> tuple[list[tuple[str, ...]], list[tuple[str, ...]]]:
    """Devuelve (solo_en_disco, solo_en_excel), ignorando desktop.ini/Thumbs.db."""
    fisico = {e.rel for e in entradas if not es_ignorado(e.nombre)}
    esperado = {r for r in esperado if not es_ignorado(r[-1])}
    return sorted(fisico - esperado), sorted(esperado - fisico)
