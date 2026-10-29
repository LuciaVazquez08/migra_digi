"""Fase 6: reportes Excel (resumen por cliente, incidencias y plan de archivos)."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .modelo import Incidencia, Plan, TipoAccion
from .preparacion import Preparacion

COLS_RESUMEN = [
    "CUIT detectado", "Ruta de origen", "Ruta de destino prevista", "Estado de migración",
    "Carpetas detectadas", "Archivos detectados", "URL a descargar", "Documentos descargados",
    "PDFs generados y validados", "Archivos convertidos", "Duplicados excluidos", "Archivos ignorados",
    "Archivos en resultado", "Incidencias", "Incidencias bloqueantes", "Motivo de exclusión",
]
COLS_INCIDENCIAS = [
    "CUIT", "Ruta", "Tipo de incidencia", "Descripción", "Acción automática intentada",
    "Resultado de la acción", "Requiere intervención manual", "Bloqueante", "Estado final del cliente",
]
COLS_INVENTARIO = ["Carpeta cliente", "Ruta completa", "Tipo", "Extensión", "Tamaño (bytes)", "Nivel",
                   "En Excel"]
COLS_PLAN = ["CUIT", "Acción", "Origen", "Destino (relativo al cliente)", "URL", "Detalle"]

MAX_FILAS_HOJA = 1_000_000


@dataclass
class Reporte:
    resumen: list[list] = field(default_factory=list)
    incidencias: list[list] = field(default_factory=list)
    plan: list[list] = field(default_factory=list)
    inventario: list[list] = field(default_factory=list)
    incluir_plan: bool = True

    def agregar_inventario(self, cuit: str, entradas, esperado: set | None) -> None:
        for e in entradas:
            self.inventario.append([
                cuit, str(e.ruta), "Carpeta" if e.es_dir else ("Enlace" if e.tamano < 0 else "Archivo"),
                "" if e.es_dir else Path(e.nombre).suffix.lower(), max(e.tamano, 0), len(e.rel) + 1,
                "" if esperado is None else ("Sí" if e.rel in esperado else "No"),
            ])

    def agregar(self, plan: Plan, estado: str, destino: Path, prep: Preparacion | None,
                motivo_extra: str = "") -> None:
        bloq = [i for i in plan.incidencias if i.bloqueante]
        motivo = "; ".join(f"{i.tipo}: {i.descripcion}" for i in bloq[:10])
        if len(bloq) > 10:
            motivo += f"; ... (+{len(bloq) - 10})"
        if motivo_extra:
            motivo = f"{motivo_extra}; {motivo}" if motivo else motivo_extra
        duplicados = plan.contar(TipoAccion.EXCLUIR_DUPLICADO) + sum(
            1 for _, m in (prep.excluidos if prep else []) if m.startswith("DUPLICADO_CONTENIDO"))
        self.resumen.append([
            plan.cuit, str(plan.ruta_origen), str(destino), estado,
            plan.carpetas, plan.archivos, plan.contar(TipoAccion.DESCARGAR),
            prep.descargados if prep else 0, prep.pdfs_validados if prep else 0,
            prep.convertidos if prep else 0, duplicados, plan.contar(TipoAccion.IGNORAR),
            len(prep.generados) if prep else 0, len(plan.incidencias), len(bloq), motivo,
        ])
        for i in plan.incidencias:
            self._incidencia(i, estado)
        if self.incluir_plan:
            for a in plan.acciones:
                self.plan.append([plan.cuit, a.tipo.value, str(a.origen),
                                  "\\".join(a.destino_rel) if a.destino_rel else "", a.url or "", a.detalle])

    def agregar_simple(self, cuit: str, ruta: Path, estado: str, motivo: str) -> None:
        self.resumen.append([cuit, str(ruta), "", estado] + [0] * 11 + [motivo])

    def _incidencia(self, i: Incidencia, estado: str) -> None:
        self.incidencias.append([i.cuit, i.ruta, i.tipo, i.descripcion, i.accion, i.resultado,
                                 "Sí" if i.requiere_manual and i.bloqueante else "No",
                                 "Sí" if i.bloqueante else "No", estado])

    def guardar(self, ruta: Path) -> Path:
        ruta.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook(write_only=True)
        hojas = [("Resumen por cliente", COLS_RESUMEN, self.resumen),
                 ("Detalle de incidencias", COLS_INCIDENCIAS, self.incidencias)]
        if self.incluir_plan:
            hojas.append(("Plan de archivos", COLS_PLAN, self.plan))
        if self.inventario:
            hojas.append(("Inventario", COLS_INVENTARIO, self.inventario))
        for nombre, cols, filas in hojas:
            for parte in range(0, max(1, len(filas)), MAX_FILAS_HOJA):
                titulo = nombre if parte == 0 else f"{nombre[:25]} {parte // MAX_FILAS_HOJA + 1}"
                ws = wb.create_sheet(titulo[:31])
                for idx, ancho in enumerate([22] * len(cols), start=1):
                    ws.column_dimensions[get_column_letter(idx)].width = ancho
                ws.freeze_panes = "A2"
                ws.append(cols)
                for fila in filas[parte:parte + MAX_FILAS_HOJA]:
                    ws.append(fila)
        tmp = ruta.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        tmp.replace(ruta)
        return ruta
