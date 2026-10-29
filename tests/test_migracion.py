"""Pruebas de punta a punta del diagnóstico, la preparación y la publicación."""
import os
from pathlib import Path

import pytest

from migracion_digicorp.config import MODO_REAL
from migracion_digicorp.estado import Estado
from migracion_digicorp.inventario import leer_referencia
from migracion_digicorp.orquestador import (COMANDO_DIAGNOSTICO, COMANDO_INVENTARIO, COMANDO_MIGRAR,
                                            Opciones, ejecutar)
from migracion_digicorp.validar_pdf import validar_pdf

from .conftest import (CUIT_A, CUIT_B, CUIT_C, RespuestaFalsa, contenido_url, cuit, escribir, huella_arbol,
                       pdf_bytes, url_s1)


def correr(entorno, comando=COMANDO_MIGRAR, real=True, rename=None, **op):
    entorno.cfg.modo = MODO_REAL if real else "simulacion"
    estado = Estado(entorno.cfg.archivo_estado)
    try:
        return ejecutar(entorno.cfg, Opciones(comando=comando, confirmar_real=real and comando == COMANDO_MIGRAR,
                                              **op),
                        referencia=leer_referencia(entorno.cfg.excel_referencia, "Arbol_Carpetas"),
                        descargador=entorno.descargador(estado), estado=estado, rename=rename)
    finally:
        estado.cerrar()


def archivos_destino(entorno, c):
    raiz = entorno.destino / c
    return sorted(str(p.relative_to(raiz)).replace(os.sep, "/") for p in raiz.rglob("*") if p.is_file())


def estado_de(entorno, c):
    e = Estado(entorno.cfg.archivo_estado)
    try:
        return e.estado_cliente(c)
    finally:
        e.cerrar()


def test_caso_normal_estructura_exacta(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    antes = huella_arbol(entorno.origen)
    res = correr(entorno)
    assert res.publicados == 1, res.contadores
    assert archivos_destino(entorno, CUIT_A) == sorted([
        "Documentacion Corpo Fija/60204441/Contrato.pdf",
        "Documentacion Corpo Fija/Documentacion Facultativa/Estatuto.pdf",
        "Documentacion Corpo Fija/Documentacion Facultativa/foto dni.pdf",   # jpg convertido
        "Documentacion Corpo Fija/Documentacion Facultativa/nota.msg",      # sin conversión posible
        "Documentacion Movil/29067459SEC/DNI 1.pdf",
        "Documentacion Movil/29067459SEC/SDS.pdf",
        "Documentacion Movil/Documentacion Facultativa/Poder.pdf",          # URL facultativa propia
    ])  # 'Copia SDS.url' excluida; placeholders vacíos y desktop.ini no se migran
    for p in (entorno.destino / CUIT_A).rglob("*.pdf"):
        assert validar_pdf(p).ok
    assert huella_arbol(entorno.origen) == antes  # el origen no se toca
    assert estado_de(entorno, CUIT_A) == "PUBLICADO"
    assert not any((entorno.cfg.dir_staging).iterdir())
    assert (entorno.cfg.dir_manifiestos / f"{CUIT_A}.json").exists()


def test_simulacion_no_publica(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    res = correr(entorno, real=False)
    assert res.contadores == {"SIMULADO_OK": 1}
    assert not (entorno.destino / CUIT_A).exists()


def test_modo_real_exige_confirmacion(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    entorno.cfg.modo = MODO_REAL
    with pytest.raises(SystemExit):
        ejecutar(entorno.cfg, Opciones(comando=COMANDO_MIGRAR, confirmar_real=False),
                 referencia=leer_referencia(entorno.cfg.excel_referencia, "Arbol_Carpetas"))


def test_diagnostico_e_inventario_no_escriben_ni_descargan(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    for comando in (COMANDO_INVENTARIO, COMANDO_DIAGNOSTICO):
        res = correr(entorno, comando=comando, real=False)
        assert res.contadores == {"APTO": 1}
    assert entorno.sesion.llamadas == []
    assert not entorno.destino.exists()


def _bloqueado(entorno, c, tipo_esperado):
    res = correr(entorno)
    assert not (entorno.destino / c).exists(), "no debe publicarse nada del cliente"
    assert estado_de(entorno, c) in ("REVISION_MANUAL", "NO_APTO")
    from openpyxl import load_workbook
    wb = load_workbook(res.reporte, read_only=True)
    tipos = {f[2] for f in wb["Detalle de incidencias"].iter_rows(min_row=2, values_only=True) if f[0] == c
             and f[7] == "Sí"}
    assert tipo_esperado in tipos, tipos
    return res


def test_sds_vacia_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    (raiz / "Documentación Móvil" / "29999999SEC").mkdir()
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "SDS_VACIA")


def test_contrato_vacio_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    (raiz / "Contratos Fija" / "60999999").mkdir()
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "CONTRATO_VACIA")


@pytest.mark.parametrize("nombre", ["2012345678", "23324234233(1)", "11111111111"])
def test_cuit_incorrecto(entorno, nombre):
    entorno.cliente_normal(nombre)
    entorno.excel()
    _bloqueado(entorno, nombre, "CUIT_INVALIDO")


def test_documento_faltante_respecto_del_excel(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel(extra=[(CUIT_A, "Contratos Fija", "60204441", "Anexo.pdf")])
    _bloqueado(entorno, CUIT_A, "DIFERENCIA_EXCEL")


def test_cliente_no_en_excel(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel(omitir={CUIT_A})
    _bloqueado(entorno, CUIT_A, "CLIENTE_NO_EN_EXCEL")


def test_url_rota_bloquea_cliente_completo_y_no_afecta_a_otros(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    entorno.cliente_normal(CUIT_B, sds="29000001SEC", contrato="60000001")
    u = url_s1("rota")
    entorno.sesion.respuestas[u] = RespuestaFalsa(404, b"")
    escribir(raiz / "Documentación Móvil" / "29067459SEC" / "SPN.url", contenido_url(u))
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "DESCARGA_FALLIDA")
    assert estado_de(entorno, CUIT_B) == "PUBLICADO"


def test_url_truncada_bloquea_sin_descargar(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / "Documentación Móvil" / "29067459SEC" / "SPN.url", contenido_url(url_s1("t")[:-9]))
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "URL_INVALIDA")


def test_pdf_corrupto_en_origen_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / f"Documentación Facultativa {CUIT_A}" / "Roto.pdf", pdf_bytes()[:120])
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "PDF_CORRUPTO")


def test_descarga_html_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    u = url_s1("login")
    entorno.sesion.respuestas[u] = RespuestaFalsa(200, b"<html><form>login</form></html>")
    escribir(raiz / "Contratos Fija" / "60204441" / "Anexo.url", contenido_url(u))
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "DESCARGA_FALLIDA")


def test_copia_sin_original_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / f"Documentación Facultativa {CUIT_A}" / "Copia Estatuto.url",
             contenido_url(entorno.sesion.pdf(url_s1("huerfana"))))
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "COPIA_SIN_ORIGINAL")


def test_duplicado_por_contenido_en_misma_carpeta(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    sds = raiz / "Documentación Móvil" / "29067459SEC"
    escribir(sds / "SDS firmada.url", (sds / "SDS.url").read_bytes())  # mismo enlace, otro nombre
    entorno.excel()
    assert correr(entorno).publicados == 1
    pdfs = [p for p in archivos_destino(entorno, CUIT_A) if "29067459SEC" in p]
    assert pdfs == ["Documentacion Movil/29067459SEC/DNI 1.pdf", "Documentacion Movil/29067459SEC/SDS firmada.pdf"]


def test_colision_de_nombres_se_renombra(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    fac = raiz / f"Documentación Facultativa {CUIT_A}"
    escribir(fac / "foto dni.pdf", pdf_bytes(4))  # colisiona con foto dni.jpg -> foto dni.pdf
    entorno.excel()
    assert correr(entorno).publicados == 1
    nombres = archivos_destino(entorno, CUIT_A)
    assert "Documentacion Corpo Fija/Documentacion Facultativa/foto dni.pdf" in nombres
    assert "Documentacion Corpo Fija/Documentacion Facultativa/foto dni (2).pdf" in nombres


def test_placeholder_con_contenido_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / "Documentación Móvil" / "Numero de SDS" / "29071685SEC" / "SDS.url",
             contenido_url(entorno.sesion.pdf(url_s1("ph"))))
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "PLACEHOLDER_CON_CONTENIDO")


def test_archivo_suelto_inferible_se_reubica(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / "Contratos Fija" / "OS 60374809.pdf", pdf_bytes())
    entorno.excel()
    assert correr(entorno).publicados == 1
    assert "Documentacion Corpo Fija/60374809/OS 60374809.pdf" in archivos_destino(entorno, CUIT_A)


def test_archivo_suelto_no_inferible_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / "Documentación Móvil" / "DNI.url", contenido_url(entorno.sesion.pdf(url_s1("suelto"))))
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "ARCHIVO_SUELTO")


@pytest.mark.parametrize("nombre,tipo", [("descarga.pdf.crdownload", "ARCHIVO_INCOMPLETO_O_ACCESO_DIRECTO"),
                                         ("acceso.lnk", "ARCHIVO_INCOMPLETO_O_ACCESO_DIRECTO"),
                                         ("SINEXTENSION", "ARCHIVO_SIN_EXTENSION"),
                                         ("Listado.xlsx", "PLANILLA_EXCEL"),
                                         ("viejo.XLS", "PLANILLA_EXCEL"),
                                         ("macro.xlsm", "PLANILLA_EXCEL"),
                                         ("export.csv", "PLANILLA_EXCEL")])
def test_extensiones_bloqueantes(entorno, nombre, tipo):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / f"Documentación Facultativa {CUIT_A}" / nombre, b"x")
    entorno.excel()
    _bloqueado(entorno, CUIT_A, tipo)


def test_excel_en_sds_bloquea(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / "Documentación Móvil" / "29067459SEC" / "control.xlsx", b"PK\x03\x04 planilla")
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "PLANILLA_EXCEL")


def test_carpeta_desconocida_y_facultativa_ajena(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    (raiz / "Excepciones 2023").mkdir()
    entorno.cliente_normal(CUIT_B)
    os.rename(entorno.origen / CUIT_B / f"Documentación Facultativa {CUIT_B}",
              entorno.origen / CUIT_B / "Documentación Facultativa 29972086")
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "CARPETA_NO_RECONOCIDA")
    assert estado_de(entorno, CUIT_B) == "REVISION_MANUAL"


def test_sds_repetida_en_dos_clientes(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.cliente_normal(CUIT_B)  # ambos usan 29067459SEC
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "SDS_EN_VARIOS_CLIENTES")
    assert estado_de(entorno, CUIT_B) == "REVISION_MANUAL"


def test_conflicto_de_destino(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    previo = escribir(entorno.destino / CUIT_A / "preexistente.txt", b"no tocar")
    res = correr(entorno)
    assert res.publicados == 0 and estado_de(entorno, CUIT_A) == "REVISION_MANUAL"
    assert previo.read_bytes() == b"no tocar"
    assert archivos_destino(entorno, CUIT_A) == ["preexistente.txt"]


def test_error_durante_publicacion(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()

    def rename_falla(a, b):
        raise OSError("disco lleno")
    res = correr(entorno, rename=rename_falla)
    assert res.publicados == 0
    assert not (entorno.destino / CUIT_A).exists()
    assert not (entorno.cfg.dir_staging / CUIT_A).exists()
    assert estado_de(entorno, CUIT_A) == "NO_APTO"
    # en la siguiente ejecución se publica normalmente
    assert correr(entorno).publicados == 1


def test_idempotencia(entorno):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    assert correr(entorno).publicados == 1
    huella = huella_arbol(entorno.destino)
    llamadas = len(entorno.sesion.llamadas)
    res = correr(entorno)
    assert res.publicados == 0 and res.contadores == {"OMITIDO_YA_PUBLICADO": 1}
    assert huella_arbol(entorno.destino) == huella
    assert len(entorno.sesion.llamadas) == llamadas


def test_failsafe_tope_de_publicaciones(entorno):
    for c in (CUIT_A, CUIT_B, CUIT_C):
        entorno.cliente_normal(c, sds=c[2:10] + "SEC", contrato=c[3:11])
    entorno.excel()
    entorno.cfg.limite_publicaciones = 2
    res = correr(entorno)
    assert res.publicados == 2 and res.motivo_fin.startswith("TOPE_DE_SEGURIDAD")
    entorno.cfg.limite_publicaciones = 999_999  # nunca supera el tope duro
    assert entorno.cfg.limite_efectivo == 2000
    res = correr(entorno)  # reanudación: publica el que faltaba
    assert res.publicados == 1


def test_parada_manual_y_reanudacion(entorno):
    for c in (CUIT_A, CUIT_B):
        entorno.cliente_normal(c, sds=c[2:10] + "SEC", contrato=c[3:11])
    entorno.excel()
    entorno.cfg.trabajo.mkdir(parents=True, exist_ok=True)
    entorno.cfg.archivo_detener.write_text("")
    res = correr(entorno)
    assert res.motivo_fin == "DETENIDO_MANUALMENTE" and res.publicados == 0
    entorno.cfg.archivo_detener.unlink()
    res = correr(entorno)
    assert res.publicados == 2


def test_cliente_sin_documentacion(entorno):
    c = cuit("2099999999")
    raiz = entorno.origen / c
    (raiz / "Documentación Móvil" / "Numero de SDS").mkdir(parents=True)
    (raiz / "Contratos Fija" / "Numero Os" / "Contrato Firmado").mkdir(parents=True)
    (raiz / f"Documentación Facultativa {c}").mkdir()
    entorno.excel()
    _bloqueado(entorno, c, "CLIENTE_SIN_DOCUMENTACION")


def test_interrupcion_durante_preparacion_no_deja_restos(entorno, monkeypatch):
    entorno.cliente_normal(CUIT_A)
    entorno.excel()
    from migracion_digicorp import preparacion
    original = preparacion._materializar
    llamadas = {"n": 0}

    def materializar_y_detener(*a, **kw):
        llamadas["n"] += 1
        if llamadas["n"] == 2:
            entorno.cfg.archivo_detener.write_text("")
        return original(*a, **kw)
    monkeypatch.setattr(preparacion, "_materializar", materializar_y_detener)
    res = correr(entorno)
    assert res.motivo_fin == "DETENIDO_MANUALMENTE" and res.publicados == 0
    assert not (entorno.destino / CUIT_A).exists()
    assert list(entorno.cfg.dir_staging.iterdir()) == []
    entorno.cfg.archivo_detener.unlink()
    monkeypatch.setattr(preparacion, "_materializar", original)
    assert correr(entorno).publicados == 1


@pytest.mark.skipif(not __import__("shutil").which("soffice"), reason="LibreOffice no instalado")
def test_conversion_office(entorno):
    from .docx_minimo import crear_docx
    raiz = entorno.cliente_normal(CUIT_A)
    crear_docx(raiz / f"Documentación Facultativa {CUIT_A}" / "Filmina.docx")
    entorno.excel()
    entorno.cfg.convertir_office = True
    assert correr(entorno).publicados == 1
    assert "Documentacion Corpo Fija/Documentacion Facultativa/Filmina.pdf" in archivos_destino(entorno, CUIT_A)


def test_cliente_solo_con_url_invalidas_no_se_reporta_sin_documentacion(entorno):
    c = cuit("2088888888")
    escribir(entorno.origen / c / "Documentación Móvil" / "29111111SEC" / "SDS.url",
             contenido_url("https://claro.my.salesforce.com/sfc/p/#f4000004kWgH/a/3n0000002gJU/xyz"))
    entorno.excel()
    res = _bloqueado(entorno, c, "URL_INVALIDA")
    from openpyxl import load_workbook
    tipos = {f[2] for f in load_workbook(res.reporte, read_only=True)["Detalle de incidencias"]
             .iter_rows(min_row=2, values_only=True)}
    assert "CLIENTE_SIN_DOCUMENTACION" not in tipos


def test_copia_sin_url_legible(entorno):
    raiz = entorno.cliente_normal(CUIT_A)
    escribir(raiz / f"Documentación Facultativa {CUIT_A}" / "Copia DNI 1.url", b"[InternetShortcut]\r\n")
    entorno.excel()
    _bloqueado(entorno, CUIT_A, "COPIA_ILEGIBLE")
